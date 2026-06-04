# -*- coding: utf-8 -*-
"""
Tạo công văn theo đúng định dạng file đã được anh Bình chỉnh sửa thủ công:
- Table header 3 hàng: [UBND/SXD | CHXHCNVN] / [Số | Ngày] / [V/v - full width]
- Body: Justify, 14pt, fi=10mm, ls=1.1, sa=2
- I/II/III: Justify, bold, sb=4, fi=10mm
- 1/2/3: Justify, bold, sb=2, fi=10mm (không gạch chân)
- a/b/c: Justify, bold, fi=10mm
- Yêu cầu: Justify, bold+italic, fi=10mm
- Kính gửi: Center, 14pt
- Nơi nhận/Ký tên: Table 2 cột
"""
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_width    = Mm(210); sec.page_height   = Mm(297)
sec.left_margin   = Mm(30);  sec.right_margin  = Mm(15)
sec.top_margin    = Mm(20);  sec.bottom_margin = Mm(20)

TNR = 'Times New Roman'

def sf(run, sz=14, bold=False, italic=False, underline=False):
    f = run.font; f.name = TNR; f.size = Pt(sz)
    f.bold = bold; f.italic = italic; f.underline = underline
    run._element.rPr.rFonts.set(qn('w:eastAsia'), TNR)

def fp(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, sb=0, sa=2, ls=1.1, fi=10):
    pf = p.paragraph_format; pf.alignment = align
    pf.space_before = Pt(sb); pf.space_after = Pt(sa)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = ls
    if fi is not None: pf.first_line_indent = Mm(fi)

def ar(p, text, sz=14, bd=False, it=False, ul=False):
    r = p.add_run(text); sf(r, sz, bd, it, ul); return r

def np(text='', align=WD_ALIGN_PARAGRAPH.JUSTIFY, sz=14, bd=False, it=False,
       sb=0, sa=2, ls=1.1, fi=10):
    p = doc.add_paragraph(); fp(p, align, sb, sa, ls, fi)
    if text: ar(p, text, sz, bd, it)
    return p

def no_border(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for n in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement(f'w:{n}'); e.set(qn('w:val'),'none'); b.append(e)
    tblPr.append(b)

def merge_row(row):
    """Merge tất cả cells trong 1 row"""
    row.cells[0].merge(row.cells[-1])

def add_bottom_border(p):
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'6')
    bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),'000000')
    pBdr.append(bot); pPr.append(pBdr)

# ══════════════════════════════════════════════════════════════
# BẢNG HEADER (3 hàng, 2 cột chính — như file gốc đã chỉnh)
# ══════════════════════════════════════════════════════════════
L = Mm(70); R = Mm(95)

hdr = doc.add_table(rows=3, cols=2)
hdr.alignment = WD_TABLE_ALIGNMENT.CENTER
no_border(hdr)
for row in hdr.rows:
    row.cells[0].width = L; row.cells[1].width = R

# Hàng 0: UBND/SXD | CỘNG HÒA/Độc lập
c0 = hdr.cell(0,0); c1 = hdr.cell(0,1)

p_ubnd = c0.paragraphs[0]
p_ubnd.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ubnd.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_ubnd.paragraph_format.line_spacing = 1.0
p_ubnd.paragraph_format.space_before = Pt(0)
p_ubnd.paragraph_format.space_after  = Pt(0)
ar(p_ubnd, 'UBND TỈNH QUẢNG TRỊ', 12, False)

p_sxd = c0.add_paragraph()
p_sxd.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sxd.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_sxd.paragraph_format.line_spacing = 1.0
p_sxd.paragraph_format.space_before = Pt(0)
p_sxd.paragraph_format.space_after  = Pt(3)
ar(p_sxd, 'SỞ XÂY DỰNG', 13, True, False, True)

p_ch = c1.paragraphs[0]
p_ch.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ch.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_ch.paragraph_format.line_spacing = 1.0
p_ch.paragraph_format.space_before = Pt(0)
p_ch.paragraph_format.space_after  = Pt(0)
ar(p_ch, 'CỘNG HOÀ XÃ HỘI CHỦ NGHĨA VIỆT NAM', 12, True)

p_dl = c1.add_paragraph()
p_dl.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_dl.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_dl.paragraph_format.line_spacing = 1.0
p_dl.paragraph_format.space_before = Pt(0)
p_dl.paragraph_format.space_after  = Pt(3)
ar(p_dl, 'Độc lập - Tự do - Hạnh phúc', 13, True, False, True)

# Hàng 1: Số | Ngày
c2 = hdr.cell(1,0); c3 = hdr.cell(1,1)

p_so = c2.paragraphs[0]
p_so.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_so.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_so.paragraph_format.line_spacing = 1.0
p_so.paragraph_format.space_before = Pt(4)
p_so.paragraph_format.space_after  = Pt(0)
ar(p_so, 'Số:        /SXD-HTKT', 13, False)

p_ng = c3.paragraphs[0]
p_ng.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ng.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_ng.paragraph_format.line_spacing = 1.0
p_ng.paragraph_format.space_before = Pt(4)
p_ng.paragraph_format.space_after  = Pt(0)
ar(p_ng, 'Quảng Trị, ngày      tháng      năm 2026', 14, False, True)

# Hàng 2: V/v trích yếu — merge full width
merge_row(hdr.rows[2])
c4 = hdr.cell(2,0)
# Xoá các đoạn trống thừa do merge tạo ra, chỉ giữ 1 đoạn
for extra in c4.paragraphs[1:]:
    p_el = extra._p
    p_el.getparent().remove(p_el)
p_vv = c4.paragraphs[0]
p_vv.alignment = WD_ALIGN_PARAGRAPH.LEFT
p_vv.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_vv.paragraph_format.line_spacing = 1.1
p_vv.paragraph_format.space_before = Pt(4)
p_vv.paragraph_format.space_after  = Pt(2)
p_vv.paragraph_format.first_line_indent = Pt(0)
ar(p_vv, 'V/v thông báo kết quả thẩm định và yêu cầu bổ sung, hoàn chỉnh hồ sơ '
    'điều chỉnh Báo cáo kinh tế - kỹ thuật dự án Tạo quỹ đất khu dân cư '
    'phía Đông đường Hà Huy Tập, tổ dân phố 6, phường Đông Sơn, tỉnh Quảng Trị', 12)

# ══════════════════════════════════════════════════════════════
# KÍNH GỬI
# ══════════════════════════════════════════════════════════════
p_kg = doc.add_paragraph()
fp(p_kg, WD_ALIGN_PARAGRAPH.CENTER, sb=2, sa=2, ls=1.1, fi=0)
ar(p_kg, 'Kính gửi: Trung tâm Phát triển quỹ đất tỉnh Quảng Trị.', 14)

# ══════════════════════════════════════════════════════════════
# HELPERS NỘI DUNG
# ══════════════════════════════════════════════════════════════
def body(text, sb=0):
    return np(text, WD_ALIGN_PARAGRAPH.JUSTIFY, 14, False, False, sb=sb, sa=2, ls=1.1, fi=10)

def h1(text, sb=4):
    """I. II. III. — Justify, bold, fi=10"""
    return np(text, WD_ALIGN_PARAGRAPH.JUSTIFY, 14, True, False, sb=sb, sa=2, ls=1.1, fi=10)

def h2(text, n, sb=2):
    """1. 2. 3. — Justify, bold, không gạch chân"""
    return np(f'{n}. {text}', WD_ALIGN_PARAGRAPH.JUSTIFY, 14, True, False, sb=sb, sa=2, ls=1.1, fi=10)

def item(label, text):
    """a) b) c) — Justify, bold cả dòng"""
    p = doc.add_paragraph(); fp(p, WD_ALIGN_PARAGRAPH.JUSTIFY, sb=0, sa=2, ls=1.1, fi=10)
    ar(p, label + ' ' + text, 14, True); return p

def yc(text):
    """Yêu cầu: — Justify, bold+italic"""
    p = doc.add_paragraph(); fp(p, WD_ALIGN_PARAGRAPH.JUSTIFY, sb=0, sa=2, ls=1.1, fi=10)
    ar(p, 'Yêu cầu: ', 14, True, True)
    ar(p, text, 14, False, False); return p

def mixed(parts, sb=0):
    p = doc.add_paragraph(); fp(p, WD_ALIGN_PARAGRAPH.JUSTIFY, sb=sb, sa=2, ls=1.1, fi=10)
    for txt, bd, it in parts: ar(p, txt, 14, bd, it)
    return p

# ══════════════════════════════════════════════════════════════
# NỘI DUNG
# ══════════════════════════════════════════════════════════════
body('Sở Xây dựng tỉnh Quảng Trị nhận được hồ sơ kèm theo Tờ trình số 109/TTr-PTQĐ '
     'của Trung tâm Phát triển quỹ đất tỉnh Quảng Trị về việc đề nghị thẩm định điều '
     'chỉnh Báo cáo kinh tế - kỹ thuật đầu tư xây dựng dự án Tạo quỹ đất khu dân cư '
     'phía Đông đường Hà Huy Tập, tổ dân phố 6, phường Đông Sơn, tỉnh Quảng Trị '
     '(sau đây gọi là dự án).')
body('Sau khi nghiên cứu hồ sơ, Sở Xây dựng có ý kiến như sau:')

h1('I. NHẬN XÉT CHUNG')
body('Hồ sơ điều chỉnh Báo cáo kinh tế - kỹ thuật do Công ty TNHH Tư vấn Thiết kế '
     'Phú Sơn lập, trình bày nội dung bổ sung các hạng mục hạ tầng kỹ thuật giai đoạn '
     '2 (vỉa hè, thoát nước mặt, cấp điện, chiếu sáng) với tổng mức đầu tư bổ sung '
     '1.739.000.000 đồng, nâng tổng mức đầu tư điều chỉnh lên 16.689.000.000 đồng. '
     'Bố cục hồ sơ cơ bản đủ thành phần theo quy định. Tuy nhiên, hồ sơ còn một số nội '
     'dung cần bổ sung, hoàn chỉnh trước khi Sở Xây dựng tiếp tục thẩm định, cụ thể như sau:')

h1('II. CÁC NỘI DUNG YÊU CẦU BỔ SUNG, HOÀN CHỈNH')

# 1
h2('Về căn cứ pháp lý lập dự toán', '1')
body('Thuyết minh dự toán xây dựng (Sheet TM, file TMDT-HA HUY TAP.xlsx) trích dẫn '
     'các văn bản quy phạm pháp luật đã hết hiệu lực, cụ thể:')

data1 = [
    ('a)', 'Nghị định số 32/2015/NĐ-CP ngày 25/3/2015 đã hết hiệu lực; đề nghị thay thế '
           'bằng Nghị định số 10/2021/NĐ-CP ngày 09/02/2021 của Chính phủ về quản lý chi '
           'phí đầu tư xây dựng.'),
    ('b)', 'Thông tư số 06/2016/TT-BXD ngày 10/3/2016 và Thông tư số 05/2016/TT-BXD ngày '
           '10/3/2016 của Bộ Xây dựng đã hết hiệu lực; đề nghị thay thế bằng Thông tư số '
           '11/2021/TT-BXD ngày 31/8/2021.'),
    ('c)', 'Thông tư số 01/2017/TT-BXD ngày 06/02/2017 về chi phí khảo sát đã hết hiệu '
           'lực; đề nghị thay thế bằng Thông tư số 11/2021/TT-BXD ngày 31/8/2021.'),
    ('d)', 'Thông tư số 09/2016/TT-BTC ngày 18/01/2016 về quyết toán dự án đã hết hiệu '
           'lực; đề nghị thay thế bằng Thông tư số 96/2021/TT-BTC ngày 11/11/2021.'),
    ('đ)', 'Thông tư số 150/2014/TT-BTC ngày 10/10/2014 về phí thẩm duyệt thiết kế PCCC '
           'đã hết hiệu lực từ ngày 01/01/2017; đề nghị thay thế bằng Thông tư số '
           '258/2016/TT-BTC ngày 11/11/2016.'),
    ('e)', 'Nghị định số 63/2014/NĐ-CP ngày 26/6/2014 về lựa chọn nhà thầu đã hết hiệu '
           'lực; đề nghị thay thế bằng Nghị định số 24/2024/NĐ-CP ngày 27/02/2024.'),
    ('g)', 'Quyết định số 79/2017/QĐ-BXD ngày 15/02/2017 về định mức chi phí quản lý dự '
           'án đã hết hiệu lực; đề nghị áp dụng Thông tư số 11/2021/TT-BXD và các văn '
           'bản hướng dẫn hiện hành.'),
    ('h)', 'Quyết định số 06/2016/QĐ-UBND ngày 29/4/2016 của UBND tỉnh Quảng Bình về '
           'biểu cước vận chuyển không còn áp dụng trên địa bàn tỉnh Quảng Trị; đề nghị '
           'cập nhật theo biểu cước hiện hành của tỉnh Quảng Trị.'),
    ('i)', 'Các công bố định mức dự toán xây dựng số 1776/BXD-VP (2007), 1129/BXD-VP '
           '(2009), 588/BXD-VP (2014) và định mức vật tư 1784/BXD-VP (2007) đã hết hiệu '
           'lực; đề nghị áp dụng Thông tư số 12/2021/TT-BXD ngày 31/8/2021.'),
]
for lbl, txt in data1: item(lbl, txt)
yc('Trung tâm Phát triển quỹ đất tỉnh Quảng Trị chỉ đạo đơn vị tư vấn rà soát, '
   'cập nhật toàn bộ căn cứ pháp lý trong thuyết minh dự toán, đảm bảo các văn bản '
   'trích dẫn còn hiệu lực tại thời điểm lập hồ sơ.')

# 2
h2('Về đơn giá vật liệu xây dựng và nhân công', '2')
body('Dự toán xây dựng được lập trên cơ sở giá vật liệu xây dựng tháng 4/2019 theo '
     'Thông báo số 1358/CB-LN ngày 03/5/2019 của liên ngành tỉnh Quảng Bình (Sheet TM, '
     'dòng cơ sở lập dự toán). Đơn giá vật liệu từ thời điểm hơn 06 năm trước, lại '
     'thuộc địa bàn tỉnh Quảng Bình, không phản ánh mặt bằng giá tại thời điểm lập '
     'hồ sơ trên địa bàn tỉnh Quảng Trị.')
yc('Trung tâm Phát triển quỹ đất tỉnh Quảng Trị chỉ đạo đơn vị tư vấn cập nhật đơn '
   'giá vật liệu xây dựng theo Thông báo giá vật liệu xây dựng mới nhất do Sở Xây '
   'dựng (hoặc liên Sở) tỉnh Quảng Trị công bố tại thời điểm lập dự toán; cập nhật '
   'đơn giá nhân công theo quy định hiện hành.')

# 3
h2('Về sai sót ngày ban hành trong Thuyết minh BCKTKT', '3')
mixed([
    ('Tại Mục I Thuyết minh Báo cáo kinh tế - kỹ thuật, Nghị định số 35/2023/NĐ-CP '
     'được ghi ngày "09/02/2023". Tuy nhiên, ngày ban hành thực tế của Nghị định này là ',
     False, False),
    ('20/06/2023', True, False),
    ('. Đề nghị đơn vị tư vấn chỉnh sửa lại ngày ban hành cho chính xác.', False, False),
])

# 4
h2('Về quy chuẩn kỹ thuật quốc gia về nước thải sinh hoạt', '4')
mixed([
    ('Thuyết minh kỹ thuật áp dụng QCVN 14:2008/BTNMT về nước thải sinh hoạt. '
     'Quy chuẩn này đã hết hiệu lực kể từ ngày ', False, False),
    ('01/9/2025', True, False),
    (', được thay thế bởi ', False, False),
    ('QCVN 14:2025/BTNMT', True, False),
    (' ban hành kèm theo Thông tư số 05/2025/TT-BTNMT ngày 28/02/2025 của Bộ Tài nguyên '
     'và Môi trường. Đề nghị đơn vị tư vấn cập nhật, áp dụng QCVN 14:2025/BTNMT '
     'trong hồ sơ thiết kế.', False, False),
])

# 5
h2('Về tiêu chuẩn kỹ thuật trong Chỉ dẫn kỹ thuật quản lý chất lượng (Mục V Thuyết minh)', '5')
body('Bảng danh mục tiêu chuẩn thí nghiệm và kiểm tra chất lượng tại Mục V Thuyết minh '
     'có nhiều tiêu chuẩn đã hết hiệu lực hoặc ghi sai số hiệu, cụ thể:')

tbl_std = doc.add_table(rows=1, cols=3)
tbl_std.style = 'Table Grid'
tbl_std.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_std.columns[0].width = Mm(12)
tbl_std.columns[1].width = Mm(55)
tbl_std.columns[2].width = Mm(103)

for cell, txt in zip(tbl_std.rows[0].cells,
    ['STT','Tiêu chuẩn trong hồ sơ','Yêu cầu cập nhật / xử lý']):
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); sf(r, 12, True)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.1

std_rows = [
    ('4','TCVN 6735:2000','Hết hiệu lực → Thay bằng TCVN 6735:2018'),
    ('5','TCVN 9385:2011','Sai số hiệu → Sửa thành TCVN 9385:2012'),
    ('6','TCVN 3118:1993','Hết hiệu lực → Thay bằng TCVN 3118:2022'),
    ('7','TCVN 3121:2003','Nhiều phần đã thay → Cập nhật theo TCVN 3121:2022'),
    ('8','TCVN 197:2002','Hết hiệu lực → Thay bằng TCVN 197-1:2014'),
    ('9','TCVN 198:1985','Hết hiệu lực → Thay bằng TCVN 198:2008'),
    ('10','TCVN 1651-1:2008','Hết hiệu lực → Thay bằng TCVN 1651-1:2018'),
    ('11','TCVN 1651-2:2008','Hết hiệu lực → Thay bằng TCVN 1651-2:2018'),
    ('14','TCVN 6260:2009','Hết hiệu lực → Thay bằng TCVN 6260:2020'),
    ('15','TCVN 4030:2003','Hết hiệu lực → Thay bằng TCVN 13605:2023'),
    ('16','TCVN 4787:2001','Hết hiệu lực → Thay bằng TCVN 4787:2009'),
    ('17','TCVN 6016:1995','Hết hiệu lực → Thay bằng TCVN 6016:2011'),
    ('18','TCVN 6017:1995','Hết hiệu lực → Thay bằng TCVN 6017:2015'),
    ('19','TCVN 4201:1995','Hết hiệu lực → Thay bằng TCVN 4201:2012'),
    ('20','TCVN 4198:1995','Hết hiệu lực → Thay bằng TCVN 4198:2014'),
    ('24','TCVN 8867:2011','Hết hiệu lực → Thay bằng TCVN 8867:2025'),
    ('3','TCXDVN 239:2006','Cập nhật → Áp dụng TCVN 14524:2025 hoặc TCVN 12252:2020'),
    ('21','22 TCN 346-06','Tiêu chuẩn ngành cũ → Làm rõ hoặc chuyển sang TCVN tương ứng'),
    ('22','22 TCN 333-06','Tiêu chuẩn ngành cũ → Xác định tiêu chuẩn đầm nén hiện hành'),
    ('12','TCVN 7572:2006','Còn áp dụng → Ghi rõ từng phần 7572-1, 7572-2... theo phép thử'),
    ('25','TCVN 8860:2011','Còn áp dụng → Ghi rõ từng phần 8860-1... theo phép thử'),
]
for s,t,x in std_rows:
    row = tbl_std.add_row().cells
    for cell, txt, al in zip(row,[s,t,x],
        [WD_ALIGN_PARAGRAPH.CENTER,WD_ALIGN_PARAGRAPH.LEFT,WD_ALIGN_PARAGRAPH.LEFT]):
        p = cell.paragraphs[0]; p.alignment = al
        r = p.add_run(txt); sf(r, 11)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.1

np('', sb=2, sa=0, fi=0)
yc('Đơn vị tư vấn rà soát, cập nhật toàn bộ danh mục tiêu chuẩn tại Mục V Thuyết '
   'minh; thay thế các tiêu chuẩn đã hết hiệu lực bằng phiên bản hiện hành; ghi rõ '
   'số hiệu đúng và từng phần cụ thể đối với bộ tiêu chuẩn nhiều phần.')

# 6
h2('Về kết cấu vỉa hè', '6')
body('Hồ sơ thiết kế quy định kết cấu lớp bê tông xi măng M150 dày 10 cm cho hạng mục '
     'vỉa hè. Đề nghị đơn vị tư vấn kiểm tra, rà soát lại cường độ bê tông vỉa hè; '
     'trường hợp giữ nguyên M150 thì phải có giải trình kỹ thuật cụ thể và căn cứ tiêu '
     'chuẩn áp dụng kèm theo.')

# 7
h2('Về tên đơn vị hành chính trong hồ sơ', '7')
mixed([
    ('Thuyết minh và một số bản vẽ kỹ thuật còn sử dụng địa danh cũ "Phường Bắc Nghĩa, '
     'thành phố Đồng Hới, tỉnh Quảng Bình" trong khi theo Tờ trình số 109/TTr-PTQĐ, '
     'địa điểm xây dựng hiện là ', False, False),
    ('"Phường Đông Sơn, tỉnh Quảng Trị"', True, False),
    ('. Đề nghị đơn vị tư vấn cập nhật thống nhất trong toàn bộ hồ sơ.', False, False),
])

# III
h1('III. THỜI HẠN BỔ SUNG HỒ SƠ')
mixed([
    ('Đề nghị Trung tâm Phát triển quỹ đất tỉnh Quảng Trị chỉ đạo đơn vị tư vấn '
     'hoàn chỉnh các nội dung nêu trên và nộp lại hồ sơ về Sở Xây dựng trong thời hạn ',
     False, False),
    ('15 ngày làm việc', True, False),
    (' kể từ ngày nhận được văn bản này. Trường hợp cần thêm thời gian, '
     'Trung tâm Phát triển quỹ đất tỉnh Quảng Trị có văn bản đề nghị gia hạn '
     'gửi Sở Xây dựng trước khi hết hạn.', False, False),
])
body('Sở Xây dựng tỉnh Quảng Trị thông báo để Trung tâm Phát triển quỹ đất tỉnh '
     'Quảng Trị biết, phối hợp thực hiện./.', sb=0)

np('', sb=6, sa=0, fi=0)

# ══════════════════════════════════════════════════════════════
# NƠI NHẬN & KÝ TÊN
# ══════════════════════════════════════════════════════════════
ky = doc.add_table(rows=1, cols=2)
ky.alignment = WD_TABLE_ALIGNMENT.CENTER
no_border(ky)
ky.cell(0,0).width = L; ky.cell(0,1).width = R

nn = ky.cell(0,0)
p0 = nn.paragraphs[0]
p0.paragraph_format.space_before = Pt(0)
p0.paragraph_format.space_after = Pt(0)
p0.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p0.paragraph_format.line_spacing = 1.1
ar(p0, 'Nơi nhận:', 12, True, True)
for txt in ['- Như trên;', '- UBND tỉnh (b/c);', '- Lưu: VT, HTKT.']:
    pn = nn.add_paragraph(txt)
    pn.paragraph_format.space_before = Pt(0)
    pn.paragraph_format.space_after = Pt(0)
    pn.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pn.paragraph_format.line_spacing = 1.1
    for r in pn.runs: sf(r, 12)

kt = ky.cell(0,1)
for txt in ['KT. GIÁM ĐỐC', 'PHÓ GIÁM ĐỐC']:
    p = kt.paragraphs[0] if txt=='KT. GIÁM ĐỐC' else kt.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.1
    ar(p, txt, 13, True)

for _ in range(4):
    pe = kt.add_paragraph()
    pe.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pe.paragraph_format.line_spacing = 1.2

p_ten = kt.add_paragraph()
p_ten.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ten.paragraph_format.space_before = Pt(0)
p_ten.paragraph_format.space_after = Pt(0)
p_ten.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
p_ten.paragraph_format.line_spacing = 1.0
ar(p_ten, 'Nguyễn Xuân Hoàng', 14, True)

# ══════════════════════════════════════════════════════════════
out = r'D:\Dropbox\Dropbox\Luu_du_lieu\Binh 2026\Tham dinh\Đường Hà Huy Tập TDP6\CV_yeu_cau_bo_sung_HA_HUY_TAP_v9.docx'
doc.save(out)
import sys; sys.stdout.buffer.write(('OK: '+out+'\n').encode('utf-8'))
