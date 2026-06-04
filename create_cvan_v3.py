# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_width    = Mm(210); sec.page_height   = Mm(297)
sec.left_margin   = Mm(30);  sec.right_margin  = Mm(15)
sec.top_margin    = Mm(20);  sec.bottom_margin = Mm(20)

def sfont(run, size=14, bold=False, italic=False, underline=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold; run.font.italic = italic; run.font.underline = underline
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

def para(text='', align=WD_ALIGN_PARAGRAPH.LEFT, size=14, bold=False, italic=False,
         underline=False, sb=0, sa=2, ls=1.1, fi=None):
    p = doc.add_paragraph(); p.alignment = align
    pf = p.paragraph_format
    pf.space_before = Pt(sb); pf.space_after = Pt(sa)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE; pf.line_spacing = ls
    if fi is not None: pf.first_line_indent = Mm(fi)
    if text:
        r = p.add_run(text); sfont(r, size, bold, italic, underline)
    return p

def run(p, text, size=14, bold=False, italic=False, underline=False):
    r = p.add_run(text); sfont(r, size, bold, italic, underline); return r

def no_border(tbl):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement('w:tblBorders')
    for n in ['top','left','bottom','right','insideH','insideV']:
        e = OxmlElement(f'w:{n}'); e.set(qn('w:val'),'none'); b.append(e)
    tblPr.append(b)

def h1(text): para(text, WD_ALIGN_PARAGRAPH.CENTER, 14, True, sb=4, sa=2)
def h2(text):
    p = para(align=WD_ALIGN_PARAGRAPH.LEFT, sb=2, sa=2, fi=7)
    run(p, text, 14, True, underline=True)

# ── HEADER ──────────────────────────────────────────────────
t1 = doc.add_table(1,2); t1.alignment = WD_TABLE_ALIGNMENT.CENTER
t1.columns[0].width = Mm(85); t1.columns[1].width = Mm(95); no_border(t1)
lc = t1.cell(0,0); lc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
p = lc.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('UBND TỈNH QUẢNG TRỊ'); sfont(r,12,True)
p2 = lc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run('SỞ XÂY DỰNG'); sfont(r2,13,True,underline=True)
rc = t1.cell(0,1); rc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
p3 = rc.paragraphs[0]; p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r3 = p3.add_run('CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM'); sfont(r3,12,True)
p4 = rc.add_paragraph(); p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
r4 = p4.add_run('Độc lập - Tự do - Hạnh phúc'); sfont(r4,13,True,underline=True)

t2 = doc.add_table(1,2); t2.alignment = WD_TABLE_ALIGNMENT.CENTER
t2.columns[0].width = Mm(85); t2.columns[1].width = Mm(95); no_border(t2)
lc2 = t2.cell(0,0); p5 = lc2.paragraphs[0]; p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
r5 = p5.add_run('Số:         /SXD-HTKT'); sfont(r5,13)
rc2 = t2.cell(0,1); p6 = rc2.paragraphs[0]; p6.alignment = WD_ALIGN_PARAGRAPH.CENTER
r6 = p6.add_run('Quảng Trị, ngày      tháng      năm 2026'); sfont(r6,14,italic=True)

para('',sb=4,sa=0)
for line in [
    'V/v thông báo kết quả thẩm định và yêu cầu bổ sung,',
    'hoàn chỉnh hồ sơ điều chỉnh Báo cáo kinh tế - kỹ thuật dự án',
    'Tạo quỹ đất khu dân cư phía Đông đường Hà Huy Tập,',
    'tổ dân phố 6, phường Đông Sơn, tỉnh Quảng Trị',
]:
    para(line, WD_ALIGN_PARAGRAPH.CENTER, 13, True, sb=0, sa=0)

para('',sb=2,sa=0)
p_kg = para(ls=1.1,sb=0,sa=2)
run(p_kg,'Kính gửi: ',14,True,italic=True)
run(p_kg,'Trung tâm Phát triển quỹ đất tỉnh Quảng Trị.',14,italic=True)
para('',sb=2,sa=0)

# Mở đầu
p0 = para(ls=1.1,sb=0,sa=2,fi=7)
run(p0,'Sở Xây dựng tỉnh Quảng Trị nhận được hồ sơ kèm theo Tờ trình số 109/TTr-PTQĐ '
    'của Trung tâm Phát triển quỹ đất tỉnh Quảng Trị về việc đề nghị thẩm định điều chỉnh '
    'Báo cáo kinh tế - kỹ thuật đầu tư xây dựng dự án Tạo quỹ đất khu dân cư phía Đông '
    'đường Hà Huy Tập, tổ dân phố 6, phường Đông Sơn, tỉnh Quảng Trị (sau đây gọi là dự án).')
p01 = para(ls=1.1,sb=0,sa=2,fi=7)
run(p01,'Sau khi nghiên cứu hồ sơ, Sở Xây dựng có ý kiến như sau:')

# I
h1('I. NHẬN XÉT CHUNG')
p_nxc = para(ls=1.1,sb=0,sa=2,fi=7)
run(p_nxc,'Hồ sơ điều chỉnh Báo cáo kinh tế - kỹ thuật do Công ty TNHH Tư vấn Thiết kế Phú Sơn lập, '
    'trình bày nội dung bổ sung các hạng mục hạ tầng kỹ thuật giai đoạn 2 (vỉa hè, thoát nước mặt, '
    'cấp điện, chiếu sáng) với tổng mức đầu tư bổ sung 1.739.000.000 đồng, nâng tổng mức đầu tư '
    'điều chỉnh lên 16.689.000.000 đồng. Bố cục hồ sơ cơ bản đủ thành phần. '
    'Tuy nhiên, hồ sơ còn một số nội dung cần bổ sung, hoàn chỉnh trước khi Sở Xây dựng tiếp tục '
    'thẩm định, cụ thể như sau:')

# II
h1('II. CÁC NỘI DUNG YÊU CẦU BỔ SUNG, HOÀN CHỈNH')

# 1
h2('1. Về căn cứ pháp lý lập dự toán (Sheet TM, file TMDT-HA HUY TAP.xlsx)')
p1a = para(ls=1.1,sb=0,sa=2,fi=7)
run(p1a,'Thuyết minh dự toán xây dựng trích dẫn các văn bản quy phạm pháp luật đã hết hiệu lực hoặc không còn áp dụng trên địa bàn tỉnh Quảng Trị, cụ thể:')

items_1 = [
    ('a)', 'Nghị định số 32/2015/NĐ-CP ngày 25/3/2015 về quản lý chi phí đầu tư xây dựng '
          'đã hết hiệu lực; đề nghị thay thế bằng Nghị định số 10/2021/NĐ-CP ngày 09/02/2021.'),
    ('b)', 'Thông tư số 06/2016/TT-BXD ngày 10/3/2016 và Thông tư số 05/2016/TT-BXD ngày 10/3/2016 '
          'của Bộ Xây dựng đã hết hiệu lực; đề nghị thay thế bằng Thông tư số 11/2021/TT-BXD ngày 31/8/2021.'),
    ('c)', 'Thông tư số 01/2017/TT-BXD ngày 06/02/2017 của Bộ Xây dựng về chi phí khảo sát '
          'đã hết hiệu lực; đề nghị thay thế bằng Thông tư số 11/2021/TT-BXD ngày 31/8/2021.'),
    ('d)', 'Thông tư số 09/2016/TT-BTC ngày 18/01/2016 của Bộ Tài chính về quyết toán dự án '
          'đã hết hiệu lực; đề nghị thay thế bằng Thông tư số 96/2021/TT-BTC ngày 11/11/2021.'),
    ('đ)', 'Thông tư số 150/2014/TT-BTC ngày 10/10/2014 của Bộ Tài chính về phí thẩm duyệt '
          'thiết kế PCCC đã hết hiệu lực từ ngày 01/01/2017; đề nghị thay thế bằng '
          'Thông tư số 258/2016/TT-BTC ngày 11/11/2016.'),
    ('e)', 'Nghị định số 63/2014/NĐ-CP ngày 26/6/2014 về lựa chọn nhà thầu đã hết hiệu lực; '
          'đề nghị thay thế bằng Nghị định số 24/2024/NĐ-CP ngày 27/02/2024.'),
    ('g)', 'Quyết định số 79/2017/QĐ-BXD ngày 15/02/2017 về định mức chi phí quản lý dự án '
          'và tư vấn đầu tư xây dựng đã hết hiệu lực; đề nghị áp dụng theo '
          'Thông tư số 11/2021/TT-BXD và các văn bản hướng dẫn hiện hành.'),
    ('h)', 'Quyết định số 06/2016/QĐ-UBND ngày 29/4/2016 của UBND tỉnh Quảng Bình về biểu '
          'cước vận chuyển không còn hiệu lực áp dụng trên địa bàn tỉnh Quảng Trị; '
          'đề nghị cập nhật theo biểu cước hiện hành của tỉnh Quảng Trị.'),
    ('i)', 'Các công bố định mức dự toán xây dựng số 1776/BXD-VP (2007), 1129/BXD-VP (2009), '
          '588/BXD-VP (2014) và định mức vật tư 1784/BXD-VP (2007) của Bộ Xây dựng đã hết '
          'hiệu lực; đề nghị áp dụng theo Thông tư số 12/2021/TT-BXD ngày 31/8/2021.'),
]
for label, content in items_1:
    p = para(ls=1.1,sb=0,sa=2,fi=7)
    run(p, label+' ', 14, True)
    run(p, content)

p1y = para(ls=1.1,sb=0,sa=2,fi=7)
run(p1y,'Yêu cầu: ',14,True,italic=True)
run(p1y,'Trung tâm Phát triển quỹ đất tỉnh Quảng Trị chỉ đạo đơn vị tư vấn rà soát, cập nhật '
    'toàn bộ căn cứ pháp lý trong thuyết minh dự toán, đảm bảo các văn bản trích dẫn '
    'còn hiệu lực tại thời điểm lập hồ sơ.')

# 2
h2('2. Về đơn giá vật liệu xây dựng và nhân công')
p2a = para(ls=1.1,sb=0,sa=2,fi=7)
run(p2a,'Dự toán xây dựng được lập trên cơ sở giá vật liệu xây dựng tháng 4/2019 theo Thông báo '
    'số 1358/CB-LN ngày 03/5/2019 của liên ngành tỉnh Quảng Bình (tại Sheet TM, dòng cơ sở lập '
    'dự toán). Đơn giá vật liệu từ thời điểm hơn 06 năm trước, lại thuộc địa bàn tỉnh Quảng Bình, '
    'không phản ánh mặt bằng giá tại thời điểm lập hồ sơ trên địa bàn tỉnh Quảng Trị.')
p2b = para(ls=1.1,sb=0,sa=2,fi=7)
run(p2b,'Yêu cầu: ',14,True,italic=True)
run(p2b,'Trung tâm Phát triển quỹ đất tỉnh Quảng Trị chỉ đạo đơn vị tư vấn cập nhật đơn giá '
    'vật liệu xây dựng theo Thông báo giá vật liệu xây dựng mới nhất do Sở Xây dựng (hoặc '
    'liên Sở) tỉnh Quảng Trị công bố tại thời điểm lập dự toán; cập nhật đơn giá nhân công '
    'theo quy định hiện hành.')

# 3
h2('3. Về sai sót ngày ban hành văn bản trong Thuyết minh BCKTKT')
p3a = para(ls=1.1,sb=0,sa=2,fi=7)
run(p3a,'Tại Mục I Căn cứ pháp lý của Thuyết minh Báo cáo kinh tế - kỹ thuật, Nghị định số '
    '35/2023/NĐ-CP được ghi ngày "09/02/2023". Tuy nhiên, Nghị định 35/2023/NĐ-CP sửa đổi, '
    'bổ sung một số điều của các Nghị định thuộc lĩnh vực quản lý nhà nước của Bộ Xây dựng '
    'có ngày ban hành thực tế là ')
run(p3a,'20/06/2023',14,True)
run(p3a,'. Đề nghị đơn vị tư vấn kiểm tra, chỉnh sửa lại ngày ban hành cho chính xác.')

# 4
h2('4. Về quy chuẩn kỹ thuật quốc gia về nước thải sinh hoạt')
p4a = para(ls=1.1,sb=0,sa=2,fi=7)
run(p4a,'Thuyết minh kỹ thuật áp dụng QCVN 14:2008/BTNMT về nước thải sinh hoạt. Quy chuẩn này '
    'đã hết hiệu lực kể từ ngày ')
run(p4a,'01/9/2025',14,True)
run(p4a,', được thay thế bởi ')
run(p4a,'QCVN 14:2025/BTNMT',14,True)
run(p4a,' ban hành kèm theo Thông tư số 05/2025/TT-BTNMT ngày 28/02/2025 của Bộ Tài nguyên '
    'và Môi trường. Đề nghị đơn vị tư vấn cập nhật, áp dụng QCVN 14:2025/BTNMT trong hồ sơ thiết kế.')

# 5
h2('5. Về tiêu chuẩn kỹ thuật trong Chỉ dẫn kỹ thuật quản lý chất lượng (Mục V Thuyết minh)')
p5a = para(ls=1.1,sb=0,sa=2,fi=7)
run(p5a,'Bảng danh mục tiêu chuẩn thí nghiệm và kiểm tra chất lượng tại Mục V Thuyết minh '
    'có nhiều tiêu chuẩn đã hết hiệu lực hoặc ghi sai số hiệu, cụ thể:')

# Table of standards
tbl_std = doc.add_table(rows=1, cols=3)
tbl_std.style = 'Table Grid'
tbl_std.alignment = WD_TABLE_ALIGNMENT.CENTER
# Set column widths
tbl_std.columns[0].width = Mm(12)
tbl_std.columns[1].width = Mm(65)
tbl_std.columns[2].width = Mm(93)

# Header
hdr = tbl_std.rows[0].cells
for cell, txt, w in zip(hdr, ['STT','Tiêu chuẩn hiện tại trong hồ sơ','Yêu cầu cập nhật'], [12,65,93]):
    cell.width = Mm(w)
    p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); sfont(r,12,True)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.1

std_data = [
    ('4',  'TCVN 6735:2000',    'Hết hiệu lực — Thay bằng TCVN 6735:2018'),
    ('5',  'TCVN 9385:2011',    'Ghi sai số hiệu — Sửa thành TCVN 9385:2012'),
    ('6',  'TCVN 3118:1993',    'Hết hiệu lực — Thay bằng TCVN 3118:2022'),
    ('7',  'TCVN 3121:2003',    'Nhiều phần đã thay thế — Cập nhật theo TCVN 3121:2022 tương ứng phép thử vữa'),
    ('8',  'TCVN 197:2002',     'Hết hiệu lực — Thay bằng TCVN 197-1:2014'),
    ('9',  'TCVN 198:1985',     'Hết hiệu lực — Thay bằng TCVN 198:2008'),
    ('10', 'TCVN 1651-1:2008',  'Hết hiệu lực — Thay bằng TCVN 1651-1:2018'),
    ('11', 'TCVN 1651-2:2008',  'Hết hiệu lực — Thay bằng TCVN 1651-2:2018'),
    ('14', 'TCVN 6260:2009',    'Hết hiệu lực — Thay bằng TCVN 6260:2020'),
    ('15', 'TCVN 4030:2003',    'Hết hiệu lực — Thay bằng TCVN 13605:2023'),
    ('16', 'TCVN 4787:2001',    'Hết hiệu lực — Thay bằng TCVN 4787:2009'),
    ('17', 'TCVN 6016:1995',    'Hết hiệu lực — Thay bằng TCVN 6016:2011'),
    ('18', 'TCVN 6017:1995',    'Hết hiệu lực — Thay bằng TCVN 6017:2015'),
    ('19', 'TCVN 4201:1995',    'Hết hiệu lực — Thay bằng TCVN 4201:2012'),
    ('20', 'TCVN 4198:1995',    'Hết hiệu lực — Thay bằng TCVN 4198:2014'),
    ('24', 'TCVN 8867:2011',    'Hết hiệu lực — Thay bằng TCVN 8867:2025'),
    ('3',  'TCXD/TCXDVN 239:2006', 'Cập nhật theo TCVN 14524:2025 hoặc TCVN 12252:2020 tùy nội dung kiểm tra'),
    ('21', '22 TCN 346-06',     'Tiêu chuẩn ngành cũ — Tư vấn làm rõ hoặc chuyển sang hệ TCVN/ASTM tương ứng'),
    ('22', '22 TCN 333-06',     'Tiêu chuẩn ngành cũ — Xác định loại vật liệu và tiêu chuẩn đầm nén tương ứng'),
    ('12', 'TCVN 7572:2006',    'Còn áp dụng — Ghi rõ từng phần TCVN 7572-1:2006, 7572-2:2006... theo phép thử cụ thể'),
    ('25', 'TCVN 8860:2011',    'Còn áp dụng — Ghi rõ từng phần TCVN 8860-1:2011... theo phép thử cụ thể'),
]

for stt, tc, xu_ly in std_data:
    row = tbl_std.add_row().cells
    for cell, txt, align in zip(row,
        [stt, tc, xu_ly],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT]):
        p = cell.paragraphs[0]; p.alignment = align
        r = p.add_run(txt); sfont(r, 11)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        p.paragraph_format.line_spacing = 1.1

para('',sb=2,sa=0)
p5y = para(ls=1.1,sb=0,sa=2,fi=7)
run(p5y,'Yêu cầu: ',14,True,italic=True)
run(p5y,'Đơn vị tư vấn rà soát, cập nhật toàn bộ danh mục tiêu chuẩn tại Mục V Thuyết minh; '
    'thay thế các tiêu chuẩn đã hết hiệu lực bằng phiên bản mới nhất; ghi rõ số hiệu đúng '
    'và từng phần cụ thể đối với bộ tiêu chuẩn nhiều phần.')

# 6
h2('6. Về kết cấu vỉa hè')
p6a = para(ls=1.1,sb=0,sa=2,fi=7)
run(p6a,'Hồ sơ thiết kế quy định kết cấu bê tông xi măng M150 dày 10 cm cho hạng mục vỉa hè. '
    'Đề nghị đơn vị tư vấn kiểm tra, rà soát lại cường độ bê tông; trường hợp giữ nguyên M150 '
    'thì phải có giải trình kỹ thuật cụ thể và căn cứ tiêu chuẩn áp dụng kèm theo.')

# 7
h2('7. Về tên đơn vị hành chính trong hồ sơ')
p7a = para(ls=1.1,sb=0,sa=2,fi=7)
run(p7a,'Thuyết minh Báo cáo kinh tế - kỹ thuật và một số bản vẽ kỹ thuật còn sử dụng địa danh '
    'cũ "Phường Bắc Nghĩa, thành phố Đồng Hới, tỉnh Quảng Bình" trong khi theo Tờ trình số '
    '109/TTr-PTQĐ, địa điểm xây dựng hiện là ')
run(p7a,'"Phường Đông Sơn, tỉnh Quảng Trị"',14,True)
run(p7a,'. Đề nghị đơn vị tư vấn rà soát, cập nhật thống nhất tên đơn vị hành chính '
    'theo địa danh hiện hành trong toàn bộ hồ sơ.')

# III
h1('III. THỜI HẠN BỔ SUNG HỒ SƠ')
p_tl = para(ls=1.1,sb=0,sa=2,fi=7)
run(p_tl,'Đề nghị Trung tâm Phát triển quỹ đất tỉnh Quảng Trị chỉ đạo đơn vị tư vấn hoàn '
    'chỉnh các nội dung nêu trên và nộp lại hồ sơ về Sở Xây dựng trong thời hạn ')
run(p_tl,'15 ngày làm việc',14,True)
run(p_tl,' kể từ ngày nhận được văn bản này. Trường hợp cần thêm thời gian, Trung tâm '
    'Phát triển quỹ đất tỉnh Quảng Trị có văn bản đề nghị gia hạn gửi Sở Xây dựng trước khi hết hạn.')
p_ket = para(ls=1.1,sb=0,sa=2,fi=7)
run(p_ket,'Sở Xây dựng tỉnh Quảng Trị thông báo để Trung tâm Phát triển quỹ đất tỉnh Quảng Trị '
    'biết, phối hợp thực hiện./.'); para('',sb=4)

# Ký tên
t3 = doc.add_table(1,2); t3.alignment = WD_TABLE_ALIGNMENT.CENTER
t3.columns[0].width = Mm(85); t3.columns[1].width = Mm(95); no_border(t3)
lc3 = t3.cell(0,0)
p_nn = lc3.paragraphs[0]
r_nn = p_nn.add_run('Nơi nhận:'); sfont(r_nn,12,True,italic=True)
for nn in ['- Như trên;','- UBND tỉnh (b/c);','- Lưu: VT, HTKT.']:
    pn = lc3.add_paragraph(nn)
    pn.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pn.paragraph_format.line_spacing = 1.1
    for r in pn.runs: sfont(r,12)
rc3 = t3.cell(0,1)
for txt, bold in [('KT. GIÁM ĐỐC',True),('PHÓ GIÁM ĐỐC',True)]:
    p = rc3.add_paragraph() if txt != 'KT. GIÁM ĐỐC' else rc3.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(txt); sfont(r,14,bold)
for _ in range(4):
    pe = rc3.add_paragraph('')
    pe.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pe.paragraph_format.line_spacing = 1.1
p_ten = rc3.add_paragraph(); p_ten.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_ten = p_ten.add_run('Nguyễn Xuân Hoàng'); sfont(r_ten,14,True)

out = r'D:\Dropbox\Dropbox\Luu_du_lieu\Binh 2026\Tham dinh\Đường Hà Huy Tập TDP6\CV_yeu_cau_bo_sung_HA_HUY_TAP_v3.docx'
doc.save(out)
import sys; sys.stdout.buffer.write(('OK: '+out+'\n').encode('utf-8'))
