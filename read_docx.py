# -*- coding: utf-8 -*-
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys

path = r'D:\Dropbox\Dropbox\Luu_du_lieu\Binh 2026\Tham dinh\Đường Hà Huy Tập TDP6\CV_yeu_cau_bo_sung_HA_HUY_TAP.docx'
doc = Document(path)
out = []

al_map = {0:'Left',1:'Center',2:'Right',3:'Justify',None:'?'}

def get_run_info(p):
    if not p.runs:
        return '?','?','?','?'
    r = p.runs[0]
    sz = round(r.font.size.pt,1) if r.font.size else 'inherit'
    bd = r.font.bold
    it = r.font.italic
    ul = r.font.underline
    return sz, bd, it, ul

def get_para_fmt(p):
    pf = p.paragraph_format
    al = al_map.get(p.alignment, str(p.alignment))
    sb = round(pf.space_before.pt,1) if pf.space_before else 0
    sa = round(pf.space_after.pt,1)  if pf.space_after  else 0
    ls = round(float(pf.line_spacing),2) if pf.line_spacing else '?'
    fi = round(pf.first_line_indent.mm,1) if pf.first_line_indent else 0
    return al, sb, sa, ls, fi

# Paragraphs (ngoài bảng)
out.append('='*80)
out.append('PARAGRAPHS (ngoai bang)')
out.append('='*80)
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    al, sb, sa, ls, fi = get_para_fmt(p)
    sz, bd, it, ul = get_run_info(p)
    label = txt[:75]
    out.append(f'[P{i:03d}] align={al:8s} sz={sz:6} bold={str(bd):5} italic={str(it):5} underline={str(ul):5} sb={sb:4} sa={sa:4} ls={ls} fi={fi}mm')
    out.append(f'       TEXT: {label}')

# Tables
out.append('')
out.append('='*80)
out.append('TABLES')
out.append('='*80)
for ti, tbl in enumerate(doc.tables):
    out.append(f'--- Table {ti} ---')
    for ri, row in enumerate(tbl.rows):
        for ci, cell in enumerate(row.cells):
            for pi, p in enumerate(cell.paragraphs):
                txt = p.text.strip()
                if not txt:
                    continue
                al, sb, sa, ls, fi = get_para_fmt(p)
                sz, bd, it, ul = get_run_info(p)
                label = txt[:60]
                out.append(f'  T{ti}R{ri:02d}C{ci} align={al:8s} sz={sz:6} bold={str(bd):5} it={str(it):5} | {label}')

sys.stdout.buffer.write(('\n'.join(out)+'\n').encode('utf-8'))
